"""Bounded-memory exports with provenance, safe spreadsheet text and atomic files."""

import asyncio
import csv
import hashlib
import os
import re
from collections.abc import AsyncIterable, Iterable
from datetime import UTC, datetime
from pathlib import Path
from uuid import UUID, uuid4

from loguru import logger
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.domain import DocumentData, LotData
from app.repository.crawl import CrawlRepository

PROCUREMENT_COLUMNS = [
    "Lot Number",
    "Announcement Number",
    "Title",
    "Short Description",
    "Full Description",
    "Required Deliverables",
    "Budget",
    "Risk",
    "Risk Color",
    "Customer",
    "Status",
    "Deadline",
    "Delivery Place",
    "Link",
]
SOURCE_COLUMNS = [
    "Official URL",
    "Downloaded Documents",
    "Verification Time",
    "Validation Result",
    "Missing Data",
    "Risk Justification",
]
_XML_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n")
_COLORS = {"LOW": "E2F0D9", "MEDIUM": "FFF2CC", "HIGH": "FCE4D6"}


def _documents(documents: Iterable[DocumentData]) -> Iterable[DocumentData]:
    for document in documents:
        yield document
        yield from _documents(document.children)


def _safe_spreadsheet(value: object) -> object:
    if not isinstance(value, str):
        return value
    value = _XML_ILLEGAL.sub("", value)
    if value.startswith(_FORMULA_PREFIXES) or value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _row(lot: LotData) -> list:
    full_text = "\n\n".join(text for text in (lot.description, lot.specification_original) if text)
    return [
        lot.lot_number,
        lot.announcement_number,
        lot.title,
        lot.specification_summary or lot.description,
        full_text,
        lot.required_deliverables,
        lot.budget,
        lot.risk.level.value if lot.risk else None,
        lot.risk.color if lot.risk else None,
        lot.customer.name if lot.customer else None,
        lot.status,
        lot.submission_deadline.isoformat() if lot.submission_deadline else None,
        lot.delivery_place or lot.work_place or lot.service_location,
        lot.source_url,
    ]


def _source_row(lot: LotData) -> list:
    validation = lot.validation
    return [
        lot.source_url,
        "\n".join(
            f"{doc.url}"
            + (f" | {doc.local_path}" if doc.local_path else "")
            + f" | {doc.parse_status}"
            for doc in _documents(lot.documents)
        ),
        validation.verified_at.isoformat() if validation else None,
        ("VALID" if validation.valid else "INVALID")
        + (": " + "; ".join(validation.reasons) if validation.reasons else "")
        if validation
        else "NOT_VALIDATED",
        "; ".join(validation.missing_data) if validation else "validation",
        "; ".join(lot.risk.reasons) if lot.risk else "NOT_ASSESSED",
    ]


def _markdown(value: object) -> str:
    # HTML escaping prevents downloaded text from becoming active inline HTML.
    import html

    result = html.escape("" if value is None else str(value), quote=False)
    return (
        result.replace("\\", "\\\\")
        .replace("|", "\\|")
        .replace("\r\n", "\n")
        .replace("\r", "\n")
        .replace("\n", "<br>")
    )


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


class Exporter:
    def __init__(self, output_folder: Path, repository: CrawlRepository | None = None) -> None:
        self.output_folder = Path(output_folder)
        self.repository = repository

    async def export(
        self,
        lots: AsyncIterable[LotData],
        format: str,
        run_id: UUID | None = None,
        filename: str | None = None,
    ) -> Path:
        """Stream an export, publish it after completion, then record its audit entry.

        Excel includes a JSONL sidecar containing the lossless original record for
        every row. A cell over Excel's 32,767 character limit points to that sidecar.
        CSV and Markdown include provenance columns; JSON preserves the full model.
        """
        format = {"xlsx": "excel", "md": "markdown"}.get(format.lower(), format.lower())
        extensions = {"excel": ".xlsx", "csv": ".csv", "json": ".json", "markdown": ".md"}
        if format not in extensions:
            raise ValueError("Export format must be excel, csv, json or markdown")
        self.output_folder.mkdir(parents=True, exist_ok=True)
        if filename is None:
            filename = f"procurements-{datetime.now(UTC):%Y%m%dT%H%M%SZ}-{uuid4().hex[:8]}{extensions[format]}"
        if Path(filename).name != filename or filename in {"", ".", ".."}:
            raise ValueError("filename must be a simple file name within the export folder")
        if Path(filename).suffix.lower() != extensions[format]:
            raise ValueError(f"Filename must end with {extensions[format]}")
        path = self.output_folder / filename
        if path.exists():
            raise FileExistsError(f"Export already exists: {path}")
        token = uuid4().hex
        staging = path.with_name(f".{path.name}.{token}.tmp")
        sidecar = path.with_suffix(".originals.jsonl")
        sidecar_staging = sidecar.with_name(f".{sidecar.name}.{token}.tmp")
        if format == "excel" and sidecar.exists():
            raise FileExistsError(f"Export sidecar already exists: {sidecar}")
        count = 0
        published_sidecar = False
        try:
            if format == "excel":
                count = await self._excel(lots, staging, sidecar_staging, sidecar.name)
            else:
                count = await self._text(lots, staging, format)
            if format == "excel":
                os.replace(sidecar_staging, sidecar)
                published_sidecar = True
            os.replace(staging, path)
            digest = await asyncio.to_thread(_sha256, path)
            if self.repository:
                await self.repository.record_export(path, format, count, digest, run_id)
        except BaseException:
            for temporary in (staging, sidecar_staging):
                temporary.unlink(missing_ok=True)
            # A published file remains reviewable if only database auditing failed.
            if published_sidecar and not path.exists():
                sidecar.unlink(missing_ok=True)
            if path.exists():
                logger.exception("Export file exists at {} but final audit failed", path)
            raise
        logger.info("Exported {} lots as {} to {} (sha256={})", count, format, path, digest)
        return path

    async def _excel(
        self, lots: AsyncIterable[LotData], path: Path, originals: Path, sidecar_name: str
    ) -> int:
        workbook = Workbook(write_only=True)
        procurements = workbook.create_sheet("Procurements")
        sources = workbook.create_sheet("Sources")
        for sheet, columns in ((procurements, PROCUREMENT_COLUMNS), (sources, SOURCE_COLUMNS)):
            sheet.freeze_panes = "A2"
            for index in range(1, len(columns) + 1):
                sheet.column_dimensions[get_column_letter(index)].width = 24 if index < 4 else 40
            cells = []
            for name in columns:
                cell = WriteOnlyCell(sheet, name)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="203864")
                cells.append(cell)
            sheet.append(cells)
        count = 0
        try:
            with originals.open("w", encoding="utf-8", newline="\n") as lossless:
                async for lot in lots:
                    count += 1
                    if count > 1_048_575:
                        raise ValueError(
                            "Excel worksheet limit reached; use JSON/CSV or smaller runs"
                        )
                    lossless.write(lot.model_dump_json() + "\n")
                    for sheet, row in ((procurements, _row(lot)), (sources, _source_row(lot))):
                        cells = []
                        for index, value in enumerate(row):
                            value = _safe_spreadsheet(value)
                            if isinstance(value, str) and len(value) > 32767:
                                marker = f"\n[Full original: {sidecar_name}, JSONL line {count}, source_id={lot.source_id}]"
                                value = value[: 32767 - len(marker)] + marker
                            cell = WriteOnlyCell(sheet, value)
                            if isinstance(value, str):
                                cell.data_type = "s"
                            if sheet is procurements and index in (7, 8) and lot.risk:
                                cell.fill = PatternFill(
                                    "solid", fgColor=_COLORS[lot.risk.level.value]
                                )
                            cells.append(cell)
                        sheet.append(cells)
                    if count % 100 == 0:
                        await asyncio.sleep(0)
                lossless.flush()
                os.fsync(lossless.fileno())
            procurements.auto_filter.ref = f"A1:N{count + 1}"
            sources.auto_filter.ref = f"A1:F{count + 1}"
            await asyncio.to_thread(workbook.save, path)
        finally:
            for sheet in workbook.worksheets:
                if not sheet.closed:
                    sheet.close()
                if sheet._writer is not None and Path(sheet._writer.out).exists():
                    sheet._writer.cleanup()
            workbook.close()
        return count

    async def _text(self, lots: AsyncIterable[LotData], path: Path, format: str) -> int:
        count = 0
        with path.open(
            "w", encoding="utf-8-sig" if format == "csv" else "utf-8", newline=""
        ) as stream:
            writer = csv.writer(stream) if format == "csv" else None
            if writer:
                writer.writerow(PROCUREMENT_COLUMNS + SOURCE_COLUMNS)
            elif format == "json":
                stream.write("[\n")
            else:
                stream.write("# Official procurement lots\n\n")
                stream.write("| " + " | ".join(PROCUREMENT_COLUMNS + SOURCE_COLUMNS) + " |\n")
                stream.write(
                    "| " + " | ".join("---" for _ in PROCUREMENT_COLUMNS + SOURCE_COLUMNS) + " |\n"
                )
            async for lot in lots:
                if writer:
                    writer.writerow(
                        [_safe_spreadsheet(value) for value in _row(lot) + _source_row(lot)]
                    )
                elif format == "json":
                    if count:
                        stream.write(",\n")
                    stream.write(lot.model_dump_json())
                else:
                    stream.write(
                        "| "
                        + " | ".join(_markdown(value) for value in _row(lot) + _source_row(lot))
                        + " |\n"
                    )
                count += 1
                if count % 100 == 0:
                    await asyncio.sleep(0)
            if format == "json":
                stream.write("\n]\n")
            stream.flush()
            os.fsync(stream.fileno())
        return count
