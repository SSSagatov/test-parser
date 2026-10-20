from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

from openpyxl import load_workbook
from pypdf import PdfWriter

from goszakup.export import export_excel
from goszakup.specifications import DocumentText, SpecificationExtractor, extract_document_text, specification_files


def _docx_bytes():
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", """<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
          <w:body><w:p><w:r><w:t>Первый пункт</w:t></w:r></w:p>
          <w:tbl><w:tr><w:tc><w:p><w:r><w:t>Значение таблицы</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
          </w:body></w:document>""")
        archive.writestr("word/header1.xml", """<w:hdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
          <w:p><w:r><w:t>Заголовок</w:t></w:r></w:p></w:hdr>""")
    return buffer.getvalue()


def test_select_only_this_lots_specification():
    own = {"id": 10, "nameRu": "Техническая спецификация", "originalName": "lot.docx",
           "filePath": "https://ows.goszakup.gov.kz/download/lot/own"}
    matching = {"id": 11, "nameRu": "Техническое задание", "originalName": "techspec_55_123.pdf",
                "filePath": "https://ows.goszakup.gov.kz/download/trd_buy/matching"}
    other = {"id": 12, "nameRu": "Техническое задание", "originalName": "techspec_55_1234.pdf",
             "filePath": "https://ows.goszakup.gov.kz/download/trd_buy/other"}
    repeated = dict(own, filePath="https://ows.goszakup.gov.kz/download/trd_buy/repeated")
    lot = {"id": 123, "Files": [own], "TrdBuy": {"Files": [matching, other, repeated]}}
    assert specification_files(lot) == [own, matching]


def test_single_shared_announcement_specification_is_labeled(tmp_path, monkeypatch):
    shared = {"id": 40, "nameRu": "Техническая спецификация", "originalName": "обучение.pdf",
              "filePath": "https://ows.goszakup.gov.kz/download/trd_buy/shared"}
    lot = {"id": 123, "Files": [], "TrdBuy": {"Files": [shared]}}
    assert specification_files(lot) == [shared]
    config = SimpleNamespace(output_dir=tmp_path, token="secret", request_timeout=5,
                             max_retries=0, max_concurrent_requests=1)
    monkeypatch.setattr(SpecificationExtractor, "_extract_one",
                        lambda self, url: DocumentText("Текст", "Извлечено"))
    result = SpecificationExtractor(config).for_lots([lot])[123]
    assert result.text.endswith("Текст")
    assert "общий файл объявления" in result.status


def test_docx_text_includes_tables_and_header():
    text = extract_document_text(_docx_bytes())
    assert "Первый пункт" in text
    assert "Значение таблицы" in text
    assert "Заголовок" in text


def test_pdf_without_text_uses_ocr_for_each_page():
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    writer.add_blank_page(width=300, height=300)
    buffer = BytesIO()
    writer.write(buffer)
    seen = []
    text = extract_document_text(buffer.getvalue(), ocr_page=lambda _data, page: f"Страница {page + 1}",
                                 diagnostics=seen)
    assert text == "Страница 1\n\nСтраница 2"
    assert seen == ["OCR", "OCR"]


def test_excel_preserves_full_long_specification_text(tmp_path):
    text = "=Начало\n" + "А" * 65000
    lot = {"id": 123, "lotNumber": "1-ЗЦП1", "TrdBuy": {"id": 20}}
    export_excel(tmp_path / "lots.xlsx", [lot], [], {}, {123: DocumentText(text, "Извлечено")})
    book = load_workbook(tmp_path / "lots.xlsx")
    sheet = book["Lots"]
    headers = [cell.value for cell in sheet[1]]
    spec_columns = [index + 1 for index, header in enumerate(headers)
                    if header.startswith("Текст технической спецификации")]
    assert len(spec_columns) == 3
    assert "".join(sheet.cell(2, index).value or "" for index in spec_columns) == text
    assert all(sheet.cell(2, index).data_type == "s" for index in spec_columns)
    assert sheet.cell(2, headers.index("Статус техспецификации") + 1).value == "Извлечено"


def test_extractor_caches_text_and_keeps_missing_lot_explicit(tmp_path, monkeypatch):
    config = SimpleNamespace(output_dir=tmp_path, token="secret", request_timeout=5,
                             max_retries=0, max_concurrent_requests=1)
    calls = []
    def download(self, url):
        calls.append(url)
        return _docx_bytes()
    monkeypatch.setattr(SpecificationExtractor, "_download", download)
    first = {"id": 1, "Files": [{"nameRu": "Техническая спецификация", "originalName": "spec.docx",
                                 "filePath": "https://ows.goszakup.gov.kz/download/lot/one"}]}
    missing = {"id": 2, "Files": []}
    result = SpecificationExtractor(config).for_lots([first, missing])
    assert "Первый пункт" in result[1].text
    assert result[2].status == "Файл техспецификации не найден"
    SpecificationExtractor(config).for_lots([first])
    assert len(calls) == 1
