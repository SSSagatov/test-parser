"""Exports retain originals, provenance and spreadsheet safety at file boundaries."""

import csv
import json
from datetime import UTC, datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.config.settings import Settings
from app.database import Database
from app.database.models import ExportHistory
from app.export import Exporter
from app.export.service import PROCUREMENT_COLUMNS, SOURCE_COLUMNS
from app.models.domain import DocumentData, LotData, RiskResult, SearchFilters, ValidationResult
from app.repository import CrawlRepository


def sample():
    now = datetime.now(UTC)
    return LotData(
        source_id="export-unit",
        source_url="https://goszakup.gov.kz/ru/lot/view/export-unit",
        title='=HYPERLINK("https://example.invalid")',
        collected_at=now,
        description="Unit | test\n<script>unsafe()</script>",
        specification_original="Официальный текст\n" * 3000,
        documents=[
            DocumentData(
                url="https://goszakup.gov.kz/unit.pdf",
                local_path=Path("unit.pdf"),
                parse_status="parsed",
            )
        ],
        validation=ValidationResult(
            valid=False, reasons=["UNIT_TEST"], missing_data=["deadline"], verified_at=now
        ),
        risk=RiskResult(level="HIGH", color="red", reasons=["Critical documents missing"]),
    )


async def stream(*lots):
    for lot in lots:
        yield lot


@pytest.mark.asyncio
async def test_excel_columns_original_sidecar_formula_safety_and_risk_color(tmp_path):
    lot = sample()
    path = await Exporter(tmp_path).export(stream(lot), "excel", filename="unit.xlsx")
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Procurements", "Sources"]
    proc, sources = workbook.worksheets
    assert [cell.value for cell in proc[1]] == PROCUREMENT_COLUMNS
    assert [cell.value for cell in sources[1]] == SOURCE_COLUMNS
    assert proc["C2"].data_type == "s"
    assert proc["C2"].value.startswith("'=HYPERLINK")
    assert len(proc["E2"].value) <= 32767
    assert "unit.originals.jsonl" in proc["E2"].value
    assert proc["H2"].fill.fgColor.rgb == "00FCE4D6"
    assert "UNIT_TEST" in sources["D2"].value
    assert sources["E2"].value == "deadline"
    restored = LotData.model_validate_json(
        path.with_suffix(".originals.jsonl").read_text(encoding="utf-8")
    )
    assert restored == lot
    workbook.close()


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "format,filename", [("csv", "unit.csv"), ("json", "unit.json"), ("markdown", "unit.md")]
)
async def test_text_formats_keep_full_text_and_provenance(tmp_path, format, filename):
    lot = sample()
    path = await Exporter(tmp_path).export(stream(lot), format, filename=filename)
    if format == "json":
        assert LotData.model_validate(json.loads(path.read_text(encoding="utf-8"))[0]) == lot
    elif format == "csv":
        with path.open(encoding="utf-8-sig", newline="") as file:
            row = next(csv.DictReader(file))
        assert row["Title"].startswith("'=")
        assert lot.specification_original in row["Full Description"]
        assert row["Official URL"] == lot.source_url
    else:
        text = path.read_text(encoding="utf-8")
        assert "&lt;script&gt;" in text
        assert "Unit \\| test" in text
        assert "Critical documents missing" in text


@pytest.mark.asyncio
@pytest.mark.parametrize("format,suffix", [("json", "json"), ("excel", "xlsx")])
async def test_failed_stream_does_not_publish_partial_files(tmp_path, format, suffix):
    async def broken():
        yield sample()
        raise RuntimeError("stream failed")

    with pytest.raises(RuntimeError, match="stream failed"):
        await Exporter(tmp_path).export(broken(), format, filename=f"unit.{suffix}")
    assert not list(tmp_path.iterdir())


@pytest.mark.asyncio
async def test_export_is_audited_and_does_not_overwrite(tmp_path):
    db = Database(
        Settings(database_url=f"sqlite+aiosqlite:///{tmp_path / 'audit.sqlite'}", _env_file=None)
    )
    try:
        await db.create_schema()
        repo = CrawlRepository(db.session_factory)
        run = await repo.start_run(SearchFilters())
        exporter = Exporter(tmp_path, repo)
        path = await exporter.export(stream(sample()), "json", run.id, "unit.json")
        async with db.session_factory() as session:
            record = await session.scalar(select(ExportHistory))
            assert record.row_count == 1
            assert record.run_id == run.id
            assert record.path == str(path)
            assert len(record.sha256) == 64
        with pytest.raises(FileExistsError):
            await exporter.export(stream(sample()), "json", filename="unit.json")
        with pytest.raises(ValueError, match="simple file name"):
            await exporter.export(stream(), "json", filename="../escape.json")
    finally:
        await db.dispose()
