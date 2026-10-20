import json

from goszakup.state import State
from goszakup.export import announcement_url, export_excel, export_json, lot_url
from openpyxl import load_workbook


def sample_lot():
    return {"id": 12345678901234567, "lotNumber": "12345678-ЗЦП1", "trdBuyId": 25,
            "nameRu": "Ноутбук", "amount": 1200.5, "plnPointKatoList": ["550000000"],
            "Plans": [{"price": 1200.5, "refUnitsCode": "796",
                       "PlansKato": [{"refKatoCode": "550000000",
                                      "fullDeliveryPlaceNameRu": "Павлодар"}]}],
            "TrdBuy": {"id": 25, "startDate": "2026-10-15 09:00:00",
                       "endDate": "2026-10-16 09:00:00", "Files": []},
            "Files": [{"id": 10, "nameRu": "Техническая спецификация", "filePath": "/file/10"}]}


def test_checkpoint_resume_and_export(tmp_path):
    path = tmp_path / "state.sqlite3"
    state = State(path, "signature", False)
    state.add_announcements([{"id": 25}], 25)
    state.add_candidates([{"id": 12345678901234567, "trdBuyId": 25}], [25])
    state.add_candidates([{"id": 12345678901234567, "trdBuyId": 25}], [25])
    assert state.get("duplicate_ids") == 1
    state.close()
    resumed = State(path, "signature", True)
    assert resumed.get("announcement_cursor") == 25
    assert len(resumed.pending_candidates(10)) == 1
    lot = sample_lot()
    resumed.add_result(lot["id"], lot, "match")
    assert resumed.count("lots") == 1
    resumed.close()
    stats = {"lots": 1}
    export_json(tmp_path / "out.json", [lot], [], stats)
    export_excel(tmp_path / "out.xlsx", [lot], [], stats)
    data = json.loads((tmp_path / "out.json").read_text(encoding="utf-8"))
    assert set(data) == {"statistics", "lots", "review_required"}
    assert data["lots"][0]["id"] == lot["id"]
    book = load_workbook(tmp_path / "out.xlsx")
    assert book.sheetnames == ["Lots", "Full_Data", "Documents", "Review_Required", "Statistics"]
    sheet = book["Lots"]
    assert sheet["A2"].value == str(lot["id"])
    assert sheet["S2"].hyperlink.target == lot_url(lot)
    assert sheet["R2"].hyperlink.target == announcement_url(lot)
    assert book["Documents"]["D2"].value == "Техническая спецификация"


def test_long_nested_value_is_split_without_loss(tmp_path):
    lot = sample_lot()
    lot["descriptionRu"] = "a" * 65000
    export_excel(tmp_path / "long.xlsx", [lot], [], {})
    book = load_workbook(tmp_path / "long.xlsx")
    sheet = book["Full_Data"]
    headers = [cell.value for cell in sheet[1]]
    pieces = [sheet.cell(2, index + 1).value for index, header in enumerate(headers)
              if header.startswith("descriptionRu__part")]
    assert "".join(pieces) == lot["descriptionRu"]


def test_resume_removes_records_from_other_or_unknown_platforms(tmp_path):
    state = State(tmp_path / "platform.sqlite3", "signature", False)
    state.add_announcements([{"id": 1, "systemId": 3}, {"id": 2, "systemId": 4}], 2)
    state.add_candidates([{"id": 11, "trdBuyId": 1, "systemId": 3},
                          {"id": 12, "trdBuyId": 2, "systemId": 4}], [1, 2])
    state.add_result(11, {"id": 11, "systemId": 3, "TrdBuy": {"systemId": 3}}, "match")
    state.add_result(12, {"id": 12, "systemId": 4, "TrdBuy": {"systemId": 4}}, "review", "other")
    state.add_result(13, {"id": 13, "systemId": 3, "TrdBuy": {"systemId": 4}}, "match")
    state.close()
    resumed = State(tmp_path / "platform.sqlite3", "signature", True)
    assert resumed.prune_platform({1, 2, 3}) == {
        "announcements": 1, "candidates": 1, "lots": 1, "review": 1}
    assert resumed.count("announcements") == 1
    assert resumed.count("candidates") == 1
    assert resumed.count("lots") == 1
    assert resumed.count("review") == 0
    resumed.close()

