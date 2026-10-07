"""Constructed parser inputs test mechanics; none represents official procurement data."""

import asyncio
import hashlib
import io
import struct
import zlib
from pathlib import Path
from stat import S_IFLNK
from zipfile import ZipFile, ZipInfo

import openpyxl
import pymupdf
import pytest
from docx import Document
from PIL import Image

from app.config import Settings
from app.crawler.documents import DocumentProcessor
from app.models.domain import DocumentData
from app.parser.base import DocumentLimitError, summarize
from app.parser.documents import DocumentParser
from app.parser.pdf.reader import parse_pdf
from app.parser.specification import attach_specification
from app.parser.xlsx.reader import parse_xlsx


@pytest.fixture
def settings(tmp_path: Path) -> Settings:
    return Settings(_env_file=None, download_folder=tmp_path / "downloads")


def constructed_word(text: str = "Constructed test paragraph") -> bytes:
    document = Document()
    document.add_paragraph(text)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Constructed table left"
    table.cell(0, 1).text = "Constructed table right"
    document.add_paragraph("Final constructed paragraph")
    document.sections[0].header.paragraphs[0].text = "Constructed header"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def local_document(tmp_path: Path, payload: bytes, name: str = "constructed.docx") -> DocumentData:
    path = tmp_path / name
    path.write_bytes(payload)
    return DocumentData(
        url="https://zakup.gov.kz/__constructed_parser_test__", name=name, local_path=path
    )


def zip_bytes(files: dict[str, bytes]) -> bytes:
    output = io.BytesIO()
    with ZipFile(output, "w") as archive:
        for name, payload in files.items():
            archive.writestr(name, payload)
    return output.getvalue()


def constructed_stored_rar(name: str, payload: bytes) -> bytes:
    """Build a minimal RAR3 stored member so archive logic needs no external binary."""
    encoded_name = name.encode("ascii")
    main = b"Rar!\x1a\x07\x00" + struct.pack("<HBHH", 0x90CF, 0x73, 0, 13) + bytes(6)
    file_fields = (
        struct.pack(
            "<LLBLLBBHL",
            len(payload),
            len(payload),
            0,
            zlib.crc32(payload),
            (40 << 25) | (1 << 21) | (1 << 16),
            20,
            0x30,
            len(encoded_name),
            0x20,
        )
        + encoded_name
    )
    unchecked_header = struct.pack("<HBHH", 0, 0x74, 0x8000, 7 + len(file_fields)) + file_fields
    checksum = zlib.crc32(unchecked_header[2:]) & 0xFFFF
    header = struct.pack("<H", checksum) + unchecked_header[2:]
    return main + header + payload


def test_word_preserves_body_table_header_and_original(settings: Settings, tmp_path: Path) -> None:
    payload = constructed_word()
    source = local_document(tmp_path, payload)
    parsed = DocumentParser(settings).parse(source)
    assert parsed.parse_status == "success"
    assert "Constructed table left\tConstructed table right" in parsed.original_text
    assert "Constructed header" in parsed.original_text
    assert parsed.original_text.index("Constructed test paragraph") < parsed.original_text.index(
        "Final constructed paragraph"
    )
    assert parsed.local_path.read_bytes() == payload
    assert source.original_text == ""


def test_opaque_download_name_detects_office_signature(settings: Settings, tmp_path: Path) -> None:
    parsed = DocumentParser(settings).parse(
        local_document(tmp_path, constructed_word(), "opaque.bin")
    )
    assert parsed.parse_status == "success"
    assert "Final constructed paragraph" in parsed.original_text


def test_spreadsheet_reads_every_sheet_hidden_and_formulas(
    settings: Settings, tmp_path: Path
) -> None:
    path = tmp_path / "constructed.xlsx"
    book = openpyxl.Workbook()
    book.active.title = "First"
    book.active.append(["Constructed heading", 7])
    book.active.append(["Calculated", "=B1*2"])
    second = book.create_sheet("Hidden test sheet")
    second.sheet_state = "hidden"
    second.append(["Constructed last sheet"])
    book.save(path)
    result = parse_xlsx(path, settings.max_spreadsheet_cells)
    assert "[Sheet: First]" in result.text
    assert "=B1*2" in result.text
    assert "[Sheet: Hidden test sheet]" in result.text
    assert "Constructed last sheet" in result.text
    with pytest.raises(DocumentLimitError, match="MAX_SPREADSHEET_CELLS"):
        parse_xlsx(path, 1)


def test_all_pdf_pages_and_page_specific_ocr(
    settings: Settings, tmp_path: Path, monkeypatch
) -> None:
    path = tmp_path / "constructed.pdf"
    image_buffer = io.BytesIO()
    Image.new("RGB", (100, 100), "white").save(image_buffer, format="PNG")
    with pymupdf.open() as document:
        document.new_page().insert_text((72, 72), "Constructed native page one.")
        scanned = document.new_page()
        scanned.insert_image(scanned.rect, stream=image_buffer.getvalue())
        document.new_page().insert_text((72, 72), "Constructed native final page.")
        document.save(path)
    requested_pages = []

    def constructed_ocr(page, _settings):
        requested_pages.append(page.number)
        return "Constructed recognition result"

    monkeypatch.setattr("app.parser.pdf.reader._ocr_page", constructed_ocr)
    result = parse_pdf(path, settings)
    assert not result.errors
    assert requested_pages == [1]
    assert "Constructed native page one." in result.text
    assert "Constructed recognition result" in result.text
    assert "Constructed native final page." in result.text


def test_ocr_failure_never_masquerades_as_complete(
    settings: Settings, tmp_path: Path, monkeypatch
) -> None:
    path = tmp_path / "constructed.pdf"
    image_buffer = io.BytesIO()
    Image.new("RGB", (100, 100), "white").save(image_buffer, format="PNG")
    with pymupdf.open() as document:
        document.new_page().insert_text((72, 72), "Constructed readable first page.")
        scanned = document.new_page()
        scanned.insert_image(scanned.rect, stream=image_buffer.getvalue())
        document.save(path)

    def failed_ocr(*_):
        raise RuntimeError("constructed unavailable OCR")

    monkeypatch.setattr("app.parser.pdf.reader._ocr_page", failed_ocr)
    result = DocumentParser(settings).parse(
        DocumentData(url="https://zakup.gov.kz/__constructed_parser_test__", local_path=path)
    )
    assert result.parse_status == "partial_errors"
    assert "Constructed readable first page." in result.original_text
    assert "Page 2: OCR failed" in result.error


def test_empty_pdf_separator_page_does_not_require_ocr(
    settings: Settings, tmp_path: Path, monkeypatch
) -> None:
    path = tmp_path / "constructed-blank-separator.pdf"
    with pymupdf.open() as document:
        document.new_page().insert_text((72, 72), "Constructed readable first page.")
        document.new_page()
        document.save(path)

    def unexpected_ocr(*_):
        raise AssertionError("A truly blank separator page needs no OCR")

    monkeypatch.setattr("app.parser.pdf.reader._ocr_page", unexpected_ocr)
    parsed = parse_pdf(path, settings)
    assert not parsed.errors
    assert "Constructed readable first page." in parsed.text


@pytest.mark.parametrize(
    "member",
    [
        "../escape.docx",
        "/absolute.docx",
        "C:/absolute.docx",
        "sub/../../escape.docx",
        "sub\\..\\escape.docx",
    ],
)
def test_zip_traversal_rejected_before_extraction(
    settings: Settings, tmp_path: Path, member: str
) -> None:
    source = local_document(tmp_path, zip_bytes({member: constructed_word()}), "constructed.zip")
    parsed = DocumentParser(settings).parse(source)
    assert parsed.parse_status == "error"
    assert "Unsafe archive member" in parsed.error
    assert not (tmp_path / "members").exists()


def test_zip_symlink_rejected(settings: Settings, tmp_path: Path) -> None:
    buffer = io.BytesIO()
    member = ZipInfo("constructed-link")
    member.create_system = 3
    member.external_attr = (S_IFLNK | 0o777) << 16
    with ZipFile(buffer, "w") as archive:
        archive.writestr(member, "../outside")
    source = local_document(tmp_path, buffer.getvalue(), "constructed.zip")
    parsed = DocumentParser(settings).parse(source)
    assert parsed.parse_status == "error"
    assert "non-regular" in parsed.error


def test_nested_archives_preserve_provenance_and_partials(
    settings: Settings, tmp_path: Path
) -> None:
    payload = zip_bytes(
        {
            "nested.zip": zip_bytes({"constructed.docx": constructed_word()}),
            "unsupported.bin": b"constructed unsupported content",
        }
    )
    source = local_document(tmp_path, payload, "constructed.zip")
    source.is_specification = True
    parsed = DocumentParser(settings).parse(source)
    assert parsed.parse_status == "partial_errors"
    assert parsed.children[0].parse_status == "success"
    leaf = parsed.children[0].children[0]
    assert leaf.url == source.url
    assert leaf.is_specification
    assert leaf.sha256 == hashlib.sha256(leaf.local_path.read_bytes()).hexdigest()
    assert "Final constructed paragraph" in parsed.original_text
    assert "unsupported.bin" in parsed.error


def test_specification_member_is_identified_in_generically_named_archive(
    settings: Settings, tmp_path: Path
) -> None:
    payload = zip_bytes({"Техническая спецификация.docx": constructed_word()})
    source = local_document(tmp_path, payload, "constructed-documents.zip")
    parsed = DocumentParser(settings).parse(source)
    assert not parsed.is_specification
    assert parsed.children[0].is_specification
    assert parsed.children[0].parse_status == "success"


def test_nested_specification_is_attached_to_lot_once(settings: Settings, tmp_path: Path) -> None:
    from datetime import UTC, datetime

    from app.models.domain import LotData

    payload = zip_bytes(
        {"nested.zip": zip_bytes({"Техническая спецификация.docx": constructed_word()})}
    )
    parsed = DocumentParser(settings).parse(
        local_document(tmp_path, payload, "constructed-documents.zip")
    )
    lot = LotData(
        source_id="constructed-specification-test",
        source_url=parsed.url,
        collected_at=datetime.now(UTC),
        documents=[parsed],
    )
    attached = attach_specification(lot, settings.summary_characters)
    assert attached.specification_original.count("Final constructed paragraph") == 1


def test_rar_stored_member_is_parsed_and_keeps_parent_provenance(
    settings: Settings, tmp_path: Path
) -> None:
    payload = constructed_stored_rar("constructed.docx", constructed_word())
    source = local_document(tmp_path, payload, "constructed.rar")
    parsed = DocumentParser(settings).parse(source)
    assert parsed.parse_status == "success", parsed.error
    assert parsed.children[0].url == source.url
    assert "Final constructed paragraph" in parsed.original_text


def test_rar_traversal_is_rejected_before_member_write(settings: Settings, tmp_path: Path) -> None:
    payload = constructed_stored_rar("../escape.docx", constructed_word())
    parsed = DocumentParser(settings).parse(local_document(tmp_path, payload, "constructed.rar"))
    assert parsed.parse_status == "error"
    assert "Unsafe archive member" in parsed.error
    assert not (tmp_path / "members").exists()


@pytest.mark.parametrize(
    "field,value,expected",
    [
        ("max_archive_depth", 0, "MAX_ARCHIVE_DEPTH"),
        ("max_archive_members", 1, "MAX_ARCHIVE_MEMBERS"),
        ("max_extracted_bytes", 2, "MAX_EXTRACTED_BYTES"),
    ],
)
def test_archive_configured_limits(
    settings: Settings, tmp_path: Path, field: str, value: int, expected: str
) -> None:
    setattr(settings, field, value)
    source = local_document(
        tmp_path, zip_bytes({"a.bin": b"abc", "b.bin": b"def"}), "constructed.zip"
    )
    parsed = DocumentParser(settings).parse(source)
    assert parsed.parse_status == "error"
    assert expected in parsed.error


def test_office_container_expansion_limit(settings: Settings, tmp_path: Path) -> None:
    settings.max_extracted_bytes = 100
    parsed = DocumentParser(settings).parse(local_document(tmp_path, constructed_word()))
    assert parsed.parse_status == "error"
    assert "MAX_EXTRACTED_BYTES" in parsed.error


def test_empty_spreadsheet_is_not_a_readable_specification(tmp_path: Path) -> None:
    path = tmp_path / "constructed-empty.xlsx"
    book = openpyxl.Workbook()
    book.save(path)
    result = parse_xlsx(path, 100)
    assert not result.text
    assert result.errors == ["Spreadsheet contains no populated cells"]


class ConstructedTransport:
    """In-memory bytes for downloader protocol tests; never contacts a portal."""

    def __init__(self, payload: bytes, wrong_hash: bool = False) -> None:
        self.payload = payload
        self.wrong_hash = wrong_hash
        self.current = 0
        self.maximum = 0
        self.calls = 0

    async def download(self, url: str, destination: Path, max_bytes: int) -> tuple[str, int]:
        self.calls += 1
        self.current += 1
        self.maximum = max(self.maximum, self.current)
        await asyncio.sleep(0)
        destination.write_bytes(self.payload)
        self.current -= 1
        return ("wrong" if self.wrong_hash else hashlib.sha256(self.payload).hexdigest()), len(
            self.payload
        )


@pytest.mark.asyncio
async def test_download_checks_hash_and_atomically_stores_original(settings: Settings) -> None:
    payload = constructed_word()
    transport = ConstructedTransport(payload)
    processor = DocumentProcessor(settings)
    result = await processor.process(
        DocumentData(
            url="https://zakup.gov.kz/__constructed_parser_test__", name="constructed.docx"
        ),
        transport,
    )
    assert result.parse_status == "success"
    assert result.local_path.read_bytes() == payload
    assert result.sha256 == hashlib.sha256(payload).hexdigest()
    assert result.downloaded_at is not None
    assert not list(settings.download_folder.glob("*.part"))


@pytest.mark.asyncio
async def test_checksum_mismatch_and_size_limit_are_errors(settings: Settings) -> None:
    document = DocumentData(url="https://zakup.gov.kz/__constructed_parser_test__")
    processor = DocumentProcessor(settings)
    result = await processor.process(
        document, ConstructedTransport(b"constructed", wrong_hash=True)
    )
    assert result.parse_status == "error"
    assert "SHA256" in result.error
    settings.max_download_bytes = 1
    result = await processor.process(document, ConstructedTransport(b"constructed"))
    assert result.parse_status == "error"
    assert "MAX_DOWNLOAD_BYTES" in result.error
    assert not list(settings.download_folder.glob("*.part"))


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "url",
    [
        "https://example.com/doc.pdf",
        "https://zakup.gov.kz.example.com/doc.pdf",
        "file:///etc/passwd",
        "https://user:pass@zakup.gov.kz/document",
        "https://zakup.gov.kz:8443/document",
    ],
)
async def test_nonofficial_attachment_rejected_before_request(settings: Settings, url: str) -> None:
    transport = ConstructedTransport(b"unused")
    result = await DocumentProcessor(settings).process(DocumentData(url=url), transport)
    assert result.parse_status == "error"
    assert transport.calls == 0


@pytest.mark.asyncio
async def test_download_concurrency_is_bounded(settings: Settings) -> None:
    settings.download_concurrency = 2
    transport = ConstructedTransport(constructed_word())
    documents = [
        DocumentData(
            url=f"https://zakup.gov.kz/__constructed_parser_test__/{index}", name="constructed.docx"
        )
        for index in range(7)
    ]
    results = await DocumentProcessor(settings).process_many(documents, transport)
    assert len(results) == 7
    assert all(result.parse_status == "success" for result in results)
    assert transport.maximum == 2


def test_summary_is_bounded_and_uses_only_original_text() -> None:
    original = "Constructed original sentence. " * 100
    summary = summarize(original, 100)
    assert len(summary) <= 100
    assert summary.endswith("…")
    assert summary[:-1] in " ".join(original.split())
