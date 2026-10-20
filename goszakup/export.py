import json
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def announcement_url(lot):
    identifier = (lot.get("TrdBuy") or {}).get("id") or lot.get("trdBuyId")
    return f"https://goszakup.gov.kz/ru/announce/index/{identifier}" if identifier else None


def lot_url(lot):
    announcement = (lot.get("TrdBuy") or {}).get("id") or lot.get("trdBuyId")
    identifier = lot.get("id")
    if announcement and identifier:
        return f"https://goszakup.gov.kz/ru/subpriceoffer/index/{announcement}/{identifier}"
    number = lot.get("lotNumber")
    return f"https://goszakup.gov.kz/ru/search/lots?filter%5Bnumber%5D={quote(str(number))}" if number else None


def _serialize(value):
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def flatten(value, prefix=""):
    result = {}
    for key, child in value.items():
        name = f"{prefix}.{key}" if prefix else key
        if isinstance(child, dict):
            result.update(flatten(child, name))
        elif isinstance(child, list):
            result[name] = _serialize(child)
        else:
            result[name] = child
    return result


def excel_safe(value, key=""):
    if isinstance(value, (dict, list)):
        value = _serialize(value)
    if value is None:
        return None
    lower_key = key.casefold()
    if lower_key.startswith("текст технической спецификации"):
        return value
    if (lower_key.endswith("id") or lower_key.startswith("id ") or
            lower_key.startswith("номер лота") or key in {"lotNumber", "numberAnno", "customerBin", "orgBin"}):
        return str(value)
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def _sheet(book, name, headers, rows, hyperlinks=()):
    sheet = book.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        sheet.append([excel_safe(value, headers[i]) for i, value in enumerate(row)])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[1].height = 32
    for index, header in enumerate(headers, 1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = min(54, max(15, len(str(header)) + 4))
    for col in hyperlinks:
        index = headers.index(col) + 1
        for row in range(2, sheet.max_row + 1):
            cell = sheet.cell(row, index)
            if isinstance(cell.value, str) and cell.value.startswith("https://"):
                cell.hyperlink = cell.value
                cell.style = "Hyperlink"
    return sheet


def _plans(lot):
    return lot.get("Plans") or []


def _places(lot):
    return [place for plan in _plans(lot) for place in plan.get("PlansKato") or []]


def _documents(lot):
    docs = [("Лот", file) for file in lot.get("Files") or []]
    docs += [("Объявление", file) for file in (lot.get("TrdBuy") or {}).get("Files") or []]
    return docs


def export_json(path: Path, lots, review, stats):
    path.parent.mkdir(parents=True, exist_ok=True)
    jsonl_path = path.with_suffix(".jsonl")
    jsonl_temp = path.with_suffix(".jsonl.tmp")
    with jsonl_temp.open("w", encoding="utf-8", newline="\n") as handle:
        for lot in lots:
            handle.write(_serialize(lot) + "\n")
    jsonl_temp.replace(jsonl_path)
    temp = path.with_suffix(".json.tmp")
    with temp.open("w", encoding="utf-8") as handle:
        handle.write('{"statistics":')
        json.dump(stats, handle, ensure_ascii=False)
        handle.write(',"lots":[')
        with jsonl_path.open("r", encoding="utf-8") as source:
            for index, line in enumerate(source):
                if index:
                    handle.write(",")
                handle.write(line.rstrip("\n"))
        handle.write('],"review_required":')
        json.dump([{"lot": lot, "reason": reason} for lot, reason in review], handle, ensure_ascii=False)
        handle.write("}")
    temp.replace(path)


def export_excel(path: Path, lots, review, stats, specifications=None):
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.remove(book.active)
    specifications = specifications or {}
    max_spec_parts = max((max(1, (len(specifications[lot["id"]].text) + 29999) // 30000)
                          for lot in lots if lot["id"] in specifications), default=1)
    spec_headers = ["Текст технической спецификации" if part == 0 else
                    f"Текст технической спецификации (продолжение {part + 1})"
                    for part in range(max_spec_parts)]
    headers = ["ID лота", "Номер лота", "Наименование", "Описание", "Сумма, тг",
               "Количество", "Цена за единицу", "Единица измерения", "Заказчик", "БИН заказчика",
               "Организатор", "БИН организатора", "Статус объявления", "Начало приёма",
               "Окончание приёма", "КАТО поставки", "Место поставки", "Объявление", "Ссылка на лот",
               *spec_headers, "Статус техспецификации"]
    rows = []
    for lot in lots:
        ann = lot.get("TrdBuy") or {}
        plan = _plans(lot)[0] if _plans(lot) else {}
        spec = specifications.get(lot["id"])
        spec_text = spec.text if spec else ""
        spec_parts = [spec_text[index * 30000:(index + 1) * 30000]
                      for index in range(max_spec_parts)]
        rows.append([lot.get("id"), lot.get("lotNumber"), lot.get("nameRu"), lot.get("descriptionRu"),
                     lot.get("amount"), lot.get("count"), plan.get("price"), plan.get("refUnitsCode"),
                     lot.get("customerNameRu"), lot.get("customerBin"), ann.get("orgNameRu"),
                     ann.get("orgBin"), (ann.get("RefBuyStatus") or {}).get("nameRu"),
                     ann.get("startDate"), ann.get("endDate"),
                     ", ".join(str(p.get("refKatoCode")) for p in _places(lot)),
                     "; ".join(str(p.get("fullDeliveryPlaceNameRu")) for p in _places(lot)),
                     announcement_url(lot), lot_url(lot), *spec_parts,
                     spec.status if spec else "Техспецификация не обработана"])
    sheet = _sheet(book, "Lots", headers, rows, ("Объявление", "Ссылка на лот"))
    for index in range(20, 20 + max_spec_parts):
        sheet.column_dimensions[get_column_letter(index)].width = 80
        for row in range(2, sheet.max_row + 1):
            cell = sheet.cell(row, index)
            cell.data_type = "s"  # Keep even leading '=' as document text, never a formula.
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in sheet.iter_rows(min_row=2):
        row[4].number_format = '#,##0.00'
        row[6].number_format = '#,##0.00'
        for index in (13, 14):
            try:
                parsed = datetime.fromisoformat(str(row[index].value).replace("Z", "+00:00"))
                row[index].value = parsed.replace(tzinfo=None)
                row[index].number_format = "dd.mm.yyyy hh:mm"
            except (ValueError, TypeError):
                pass

    flat = [flatten(lot) for lot in lots]
    keys = sorted({key for item in flat for key in item})
    # Excel limits one cell to 32,767 characters. Split long JSON values into numbered columns.
    widths = {key: max(1, max(((len(str(item.get(key) or "")) + 29999) // 30000 for item in flat), default=1))
              for key in keys}
    full_headers = [f"{key}__part{part+1}" if widths[key] > 1 else key
                    for key in keys for part in range(widths[key])]
    full_rows = []
    for item in flat:
        row = []
        for key in keys:
            value = item.get(key)
            if value is None:
                row.extend([None] * widths[key])
            else:
                text = str(value) if isinstance(value, (dict, list, str)) else value
                if isinstance(text, str):
                    row.extend([text[i*30000:(i+1)*30000] for i in range(widths[key])])
                else:
                    row.extend([text] + [None] * (widths[key] - 1))
        full_rows.append(row)
    _sheet(book, "Full_Data", full_headers or ["Нет данных"], full_rows)

    doc_rows = []
    for lot in lots:
        for source, doc in _documents(lot):
            doc_rows.append([lot.get("id"), source, doc.get("id"), doc.get("nameRu"),
                             doc.get("originalName"), doc.get("filePath")])
    _sheet(book, "Documents", ["ID лота", "Источник", "ID документа", "Тип документа",
                               "Имя файла", "Путь API"], doc_rows, ("Путь API",))
    _sheet(book, "Review_Required", ["ID лота", "Номер лота", "Причина", "КАТО", "Ссылка на лот"],
           [[lot.get("id"), lot.get("lotNumber"), reason,
             _serialize(lot.get("plnPointKatoList")), lot_url(lot)] for lot, reason in review],
           ("Ссылка на лот",))
    _sheet(book, "Statistics", ["Показатель", "Значение"],
           [[key, value] for key, value in stats.items()])
    temp = path.with_suffix(".xlsx.tmp")
    book.save(temp)
    temp.replace(path)

